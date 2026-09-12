import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def build_docx():
    base_dir = r"c:\Users\USER\Downloads\VaultBridge"
    md_path = os.path.join(base_dir, "docs", "prompt-and-context-strategies.md")
    docx_path1 = os.path.join(base_dir, "docs", "prompt-and-context-strategies.docx")
    docx_path2 = os.path.join(base_dir, "PROMPT_AND_CONTEXT_STRATEGIES.docx")

    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    title_color = RGBColor(15, 23, 42)       # Slate 900
    primary_color = RGBColor(79, 70, 229)    # Indigo 600
    h2_color = RGBColor(13, 148, 136)        # Teal 600
    body_color = RGBColor(30, 41, 59)        # Slate 800
    code_bg = "F1F5F9"                       # Slate 100

    i = 0
    in_code_block = False
    code_lines = []

    while i < len(lines):
        line = lines[i].rstrip("\r\n")

        if line.startswith("```"):
            if in_code_block:
                in_code_block = False
                code_text = "\n".join(code_lines)
                table = doc.add_table(rows=1, cols=1)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                table.autofit = False
                cell = table.cell(0, 0)
                set_cell_background(cell, code_bg)
                set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
                cp = cell.paragraphs[0]
                cp.paragraph_format.space_before = Pt(2)
                cp.paragraph_format.space_after = Pt(2)
                c_run = cp.add_run(code_text)
                c_run.font.name = "Consolas"
                c_run.font.size = Pt(8.5)
                c_run.font.color.rgb = RGBColor(30, 41, 59)
                doc.add_paragraph().paragraph_format.space_after = Pt(6)
                code_lines = []
            else:
                in_code_block = True
                code_lines = []
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        if line.strip() == "---":
            i += 1
            continue

        if not line.strip():
            i += 1
            continue

        if line.strip().startswith("|") and line.strip().endswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
                table_lines.append(lines[i].strip())
                i += 1

            if len(table_lines) >= 2:
                parsed_rows = []
                for tline in table_lines:
                    cols = [c.strip() for c in tline.split("|")[1:-1]]
                    if all(re.match(r"^:?-+:?$", col) for col in cols if col):
                        continue
                    parsed_rows.append(cols)

                if parsed_rows:
                    num_cols = max(len(r) for r in parsed_rows)
                    tbl = doc.add_table(rows=len(parsed_rows), cols=num_cols)
                    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                    tbl.style = 'Table Grid'

                    for row_idx, rdata in enumerate(parsed_rows):
                        is_header = (row_idx == 0)
                        bg = "4F46E5" if is_header else ("F8FAFC" if row_idx % 2 == 1 else "FFFFFF")
                        for col_idx in range(num_cols):
                            cell = tbl.cell(row_idx, col_idx)
                            text_val = rdata[col_idx] if col_idx < len(rdata) else ""
                            clean_text = re.sub(r"\[([^\]]+)\]\([^\)]+\)", r"\1", text_val)
                            clean_text = clean_text.replace("`", "").replace("**", "")
                            
                            p = cell.paragraphs[0]
                            p.paragraph_format.space_before = Pt(3)
                            p.paragraph_format.space_after = Pt(3)
                            run = p.add_run(clean_text)
                            run.font.size = Pt(8.5 if not is_header else 9.0)
                            if is_header:
                                run.font.bold = True
                                run.font.color.rgb = RGBColor(255, 255, 255)
                            else:
                                run.font.color.rgb = body_color
                            set_cell_background(cell, bg)
                            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)

                    doc.add_paragraph().paragraph_format.space_after = Pt(6)
            continue

        if line.startswith("# "):
            h = doc.add_heading(line[2:], level=0)
            h.runs[0].font.size = Pt(22)
            h.runs[0].font.bold = True
            h.runs[0].font.color.rgb = title_color
            i += 1
            continue
        elif line.startswith("## "):
            h = doc.add_heading(line[3:], level=1)
            h.runs[0].font.size = Pt(15)
            h.runs[0].font.bold = True
            h.runs[0].font.color.rgb = primary_color
            i += 1
            continue
        elif line.startswith("### "):
            h = doc.add_heading(line[4:], level=2)
            h.runs[0].font.size = Pt(12)
            h.runs[0].font.bold = True
            h.runs[0].font.color.rgb = h2_color
            i += 1
            continue
        elif line.startswith("#### "):
            h = doc.add_heading(line[5:], level=3)
            h.runs[0].font.size = Pt(10.5)
            h.runs[0].font.bold = True
            h.runs[0].font.color.rgb = body_color
            i += 1
            continue

        if line.startswith("> "):
            bq_text = line[2:]
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.space_after = Pt(6)
            parts = re.split(r"(\*\*[^*]+\*\*)", bq_text)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.bold = True
                    r.font.color.rgb = primary_color
                else:
                    r = p.add_run(part)
                    r.italic = True
                    r.font.color.rgb = RGBColor(71, 85, 105)
            i += 1
            continue

        if line.startswith("- ") or line.startswith("* ") or re.match(r"^\d+\.\s", line):
            is_num = bool(re.match(r"^\d+\.\s", line))
            item_text = re.sub(r"^(-\s|\*\s|\d+\.\s)", "", line)
            
            p = doc.add_paragraph(style='List Number' if is_num else 'List Bullet')
            p.paragraph_format.space_after = Pt(3)
            
            parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", item_text)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.bold = True
                    r.font.color.rgb = body_color
                elif part.startswith("`") and part.endswith("`"):
                    r = p.add_run(part[1:-1])
                    r.font.name = "Consolas"
                    r.font.size = Pt(9)
                    r.font.color.rgb = RGBColor(185, 28, 28)
                else:
                    clean_str = re.sub(r"\[([^\]]+)\]\([^\)]+\)", r"\1", part)
                    r = p.add_run(clean_str)
                    r.font.color.rgb = body_color
            i += 1
            continue

        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", line)
        for part in parts:
            if part.startswith("**") and part.endswith("**"):
                r = p.add_run(part[2:-2])
                r.bold = True
                r.font.color.rgb = body_color
            elif part.startswith("`") and part.endswith("`"):
                r = p.add_run(part[1:-1])
                r.font.name = "Consolas"
                r.font.size = Pt(9)
                r.font.color.rgb = RGBColor(185, 28, 28)
            else:
                clean_str = re.sub(r"\[([^\]]+)\]\([^\)]+\)", r"\1", part)
                r = p.add_run(clean_str)
                r.font.color.rgb = body_color
        i += 1

    doc.save(docx_path1)
    doc.save(docx_path2)
    print(f"Generated {docx_path1}")
    print(f"Generated {docx_path2}")

if __name__ == "__main__":
    build_docx()
